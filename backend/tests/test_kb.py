"""知识库模块测试（mock 模型调用）。"""

import pytest
from unittest.mock import patch, AsyncMock


class TestDocumentUpload:
    def test_upload_unauthorized(self, client):
        resp = client.post("/api/v1/kb/upload", files={"file": ("test.txt", b"hello")})
        assert resp.status_code == 401

    def test_upload_invalid_format(self, client, user_token):
        resp = client.post(
            "/api/v1/kb/upload",
            files={"file": ("test.exe", b"malicious")},
            headers={"Authorization": f"Bearer {user_token}"},
        )
        assert resp.status_code == 400
        assert "不支持" in resp.json()["detail"]

    def test_upload_file_too_large(self, client, user_token):
        big_content = b"x" * (21 * 1024 * 1024)  # 21MB
        resp = client.post(
            "/api/v1/kb/upload",
            files={"file": ("big.pdf", big_content)},
            headers={"Authorization": f"Bearer {user_token}"},
        )
        assert resp.status_code == 400
        assert "大小超过" in resp.json()["detail"]

    @patch("app.services.kb_service._vectorize_document")
    def test_upload_path_traversal_filename(self, mock_vec, client, user_token):
        """测试路径遍历文件名被安全处理。"""
        resp = client.post(
            "/api/v1/kb/upload",
            files={"file": ("../../../etc/passwd.txt", b"malicious content")},
            headers={"Authorization": f"Bearer {user_token}"},
        )
        # 应该成功上传但文件名被清理
        assert resp.status_code == 201, resp.json()
        data = resp.json()
        # 文件名不应包含路径字符
        assert ".." not in data["filename"]
        assert data["filename"] != "../../../etc/passwd.txt"

    @patch("app.services.kb_service._vectorize_document")
    def test_upload_and_list_documents(self, mock_vec, client, user_token):
        # 上传文件
        resp = client.post(
            "/api/v1/kb/upload",
            files={"file": ("policy.txt", "奖学金申请条件：绩点3.0以上，无违纪记录。".encode("utf-8"))},
            headers={"Authorization": f"Bearer {user_token}"},
        )
        assert resp.status_code == 201, resp.json()

        # 查看列表
        resp = client.get("/api/v1/kb/documents", headers={"Authorization": f"Bearer {user_token}"})
        assert resp.status_code == 200
        assert resp.json()["total"] >= 1

    @patch("app.services.kb_service._vectorize_document")
    def test_upload_with_category(self, mock_vec, client, user_token):
        resp = client.post(
            "/api/v1/kb/upload",
            files={"file": ("policy.txt", b"test")},
            data={"category": "资助"},
            headers={"Authorization": f"Bearer {user_token}"},
        )
        assert resp.status_code == 201, resp.json()
        assert resp.json()["category"] == "资助"

    @patch("app.services.kb_service._vectorize_document")
    def test_upload_default_category(self, mock_vec, client, user_token):
        resp = client.post(
            "/api/v1/kb/upload",
            files={"file": ("policy.txt", b"test")},
            headers={"Authorization": f"Bearer {user_token}"},
        )
        assert resp.status_code == 201, resp.json()
        assert resp.json()["category"] == "综合"

    def test_extract_text_docx(self):
        import os
        import tempfile
        import docx
        from app.services import kb_service

        d = docx.Document()
        d.add_paragraph("第一条：国家助学金名额比例为在校生的 20%")
        d.add_paragraph("第二条：每年 10 月 15 日前完成评定并公示")
        table = d.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "类目"
        table.rows[0].cells[1].text = "资助"

        path = tempfile.mktemp(suffix=".docx")
        d.save(path)
        try:
            text = kb_service._extract_text(path, "docx")
            assert "国家助学金名额比例" in text
            assert "10 月 15 日" in text
            assert "资助" in text  # 表格内容也被提取
        finally:
            os.unlink(path)

    @patch("app.services.kb_service._vectorize_document")
    def test_delete_own_document(self, mock_vec, client, user_token):
        # 上传
        client.post(
            "/api/v1/kb/upload",
            files={"file": ("delete_test.txt", b"test")},
            headers={"Authorization": f"Bearer {user_token}"},
        )
        # 获取列表
        resp = client.get("/api/v1/kb/documents", headers={"Authorization": f"Bearer {user_token}"})
        docs = resp.json()["documents"]
        assert len(docs) > 0

        # 删除
        doc_id = docs[0]["id"]
        resp = client.delete(f"/api/v1/kb/documents/{doc_id}", headers={"Authorization": f"Bearer {user_token}"})
        assert resp.status_code == 200

    @patch("app.services.kb_service._vectorize_document")
    def test_delete_others_document_forbidden(self, mock_vec, client, user_token, admin_token):
        # admin 上传
        client.post(
            "/api/v1/kb/upload",
            files={"file": ("admin_doc.txt", b"admin content")},
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        # user 尝试删除 admin 的文档应失败
        resp = client.delete("/api/v1/kb/documents/1", headers={"Authorization": f"Bearer {user_token}"})
        assert resp.status_code in (403, 404)  # forbidden or not found


class TestChatPrompt:
    def test_rag_prompt_build(self):
        """测试 RAG Prompt 构建（纯函数，不需要 mock 模型）。"""
        from app.services.chat_service import _build_documents_str

        chunks = [
            {"doc_name": "学生手册.pdf", "text": "奖学金评定标准：绩点≥3.0", "distance": 0.15},
            {"doc_name": "资助政策.pdf", "text": "家庭经济困难学生可申请助学金", "distance": 0.22},
        ]
        result = _build_documents_str(chunks)
        assert "学生手册.pdf" in result
        assert "奖学金评定标准" in result
        assert "资助政策.pdf" in result
        assert "[1]" in result
        assert "[2]" in result

    def test_chunk_text(self):
        """测试文档分段函数（纯函数）。"""
        from app.services.kb_service import _chunk_text

        text = "段落一内容。\n\n段落二内容。"
        chunks = _chunk_text(text, chunk_size=1000)
        assert len(chunks) == 1
        assert "段落一" in chunks[0]

    def test_safe_filename(self):
        """测试安全文件名处理。"""
        from app.services.kb_service import _safe_filename

        assert ".." not in _safe_filename("../../../etc/passwd")
        assert _safe_filename("test.pdf") == "test.pdf"
        # 空或隐藏文件名应生成随机名
        result = _safe_filename("")
        assert result and not result.startswith(".")