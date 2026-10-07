"""通知变材料服务：一句话 → 六件套材料。"""

import io
import json
from sqlalchemy.orm import Session
from docx import Document
from app.models.notice_task import NoticeTask
from app.models.operation_log import OperationLog
from app.services.llm import call_qwen
from app.services.prompts.notice import NOTICE_SYSTEM_PROMPT, NOTICE_USER_PROMPT

# 四个 AI 生成段落：(结果字段 key, 段落标题)
SECTION_KEYS = [
    ("notice_formal", "通知-正式版"),
    ("notice_group", "通知-群发版"),
    ("notice_parent", "通知-家长版"),
    ("meeting_plan", "班会方案"),
]


def parse_sections(raw: str) -> dict:
    """宽容解析：按【】/## 标记切分，兼容常见变体；失败时正文归入正式版兜底。"""
    result = {key: "" for key, _ in SECTION_KEYS}

    positions: list[tuple[int, str, int]] = []
    for key, label in SECTION_KEYS:
        for marker in (f"【{label}】", f"## {label}", f"# {label}"):
            idx = raw.find(marker)
            if idx >= 0:
                positions.append((idx, key, len(marker)))
                break

    if not positions:
        result["notice_formal"] = raw.strip()
        return result

    positions.sort(key=lambda p: p[0])
    for i, (idx, key, mlen) in enumerate(positions):
        end = positions[i + 1][0] if i + 1 < len(positions) else len(raw)
        result[key] = raw[idx + mlen : end].strip()

    return result


def _build_signin_sheet(theme: str, event_time: str, place: str) -> str:
    lines = [
        f"# {theme} · 签到表",
        f"时间：{event_time}　地点：{place}",
        "",
        "| 序号 | 姓名 | 学号 | 班级 | 签到时间 | 备注 |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for i in range(1, 11):
        lines.append(f"| {i} |  |  |  |  |  |")
    return "\n".join(lines)


def _build_minutes_template(theme: str, event_time: str, place: str) -> str:
    return f"""# {theme} · 班会纪要（模板）

- 会议时间：{event_time}
- 会议地点：{place}
- 主持人：
- 参会人员：
- 记录人：

## 一、会议主题


## 二、主要内容
1.
2.
3.

## 三、学生反馈与讨论


## 四、后续跟进事项
| 事项 | 责任人 | 截止时间 |
| --- | --- | --- |
|  |  |  |

## 五、备注
"""


def generate_notice_package(
    db: Session,
    user_id: int,
    theme: str,
    event_time: str = "",
    place: str = "",
    audience: str = "",
) -> NoticeTask:
    """生成六件套并持久化任务。同步执行（MVP）。"""
    task = NoticeTask(
        user_id=user_id,
        theme=theme,
        event_time=event_time,
        place=place,
        audience=audience,
        status="generating",
    )
    db.add(task)
    db.commit()
    db.refresh(task)

    try:
        t = event_time or "待定"
        p = place or "待定"
        a = audience or "全体学生"
        user_prompt = NOTICE_USER_PROMPT.format(theme=theme, event_time=t, place=p, audience=a)

        raw = call_qwen(
            [
                {"role": "system", "content": NOTICE_SYSTEM_PROMPT},
                {"role": "user", "content": user_prompt},
            ],
            temperature=0.7,
        )

        sections = parse_sections(raw)
        # 代码确定性生成的两件套
        sections["signin_sheet"] = _build_signin_sheet(theme, t, p)
        sections["minutes_template"] = _build_minutes_template(theme, t, p)

        task.result_json = json.dumps(sections, ensure_ascii=False)
        task.status = "done"
        db.commit()
        db.refresh(task)

        log = OperationLog(
            user_id=user_id,
            action="generate_notice",
            detail=f"通知变材料: {theme}",
        )
        db.add(log)
        db.commit()
        return task

    except Exception as e:
        task.status = "failed"
        task.error_msg = str(e)[:500]
        db.commit()
        raise


def get_task_result(task: NoticeTask) -> dict:
    """解析任务结果 JSON，损坏时返回空结构。"""
    try:
        return json.loads(task.result_json or "{}")
    except json.JSONDecodeError:
        return {}


def parse_student_excel(file_bytes: bytes) -> list[dict]:
    """解析班级学生名单 Excel，返回 [{name, student_id, class_name}]。"""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    def col_index(*names: str) -> int:
        for n in names:
            for i, h in enumerate(headers):
                if h == n:
                    return i
        return -1

    i_name = col_index("姓名", "学生姓名", "name")
    i_id = col_index("学号", "student_id", "studentId")
    i_class = col_index("班级", "班", "class_name", "class")

    if i_name < 0:
        raise ValueError("Excel 里找不到「姓名」列")

    def cell(row, i):
        return row[i] if 0 <= i < len(row) and row[i] is not None else ""

    students = []
    for row in rows[1:]:
        if not row or not any(row):
            continue
        name = str(cell(row, i_name)).strip()
        if not name:
            continue
        students.append({
            "name": name,
            "student_id": str(cell(row, i_id)).strip(),
            "class_name": str(cell(row, i_class)).strip(),
        })
    return students


def build_signin_sheet_with_students(
    students: list[dict], theme: str = "", event_time: str = "", place: str = ""
) -> str:
    """按真实学生名单生成签到表 Markdown。"""
    lines = [
        f"# {theme or '班会'} · 签到表",
        f"时间：{event_time or '待定'}　地点：{place or '待定'}",
        "",
        "| 序号 | 姓名 | 学号 | 班级 | 签到时间 | 备注 |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for i, s in enumerate(students, 1):
        lines.append(f"| {i} | {s['name']} | {s['student_id']} | {s['class_name']} |  |  |")
    return "\n".join(lines)


def export_signin_docx(
    students: list[dict], theme: str = "", event_time: str = "", place: str = ""
) -> bytes:
    """把真实名单签到表导出为 Word。"""
    doc = Document()
    doc.styles["Normal"].font.name = "SimSun"
    title = doc.add_heading(f"{theme or '班会'} · 签到表", level=0)
    title.alignment = 1
    meta = doc.add_paragraph(f"时间：{event_time or '待定'}　地点：{place or '待定'}")
    meta.runs[0].font.size = 110000

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for j, h in enumerate(["序号", "姓名", "学号", "班级", "签到时间", "备注"]):
        table.rows[0].cells[j].text = h
    for i, s in enumerate(students, 1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = s["name"]
        cells[2].text = s["student_id"]
        cells[3].text = s["class_name"]
        cells[4].text = ""
        cells[5].text = ""

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_notice_docx(task: NoticeTask) -> bytes:
    """把六件套导出为单个 .docx。"""
    data = get_task_result(task)

    doc = Document()
    doc.styles["Normal"].font.name = "SimSun"
    title = doc.add_heading(f"通知变材料 · {task.theme}", level=0)
    title.alignment = 1

    meta = doc.add_paragraph(f"时间：{task.event_time or '待定'}　地点：{task.place or '待定'}　对象：{task.audience or '全体学生'}")
    meta.runs[0].font.size = 110000

    def add_section(heading: str, content: str):
        doc.add_heading(heading, level=1)
        for line in content.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("|") and line.strip().endswith("|"):
                p = doc.add_paragraph(line)
                p.runs[0].font.name = "Courier New"
            elif line.startswith("**") and line.endswith("**"):
                p = doc.add_paragraph()
                p.add_run(line.strip("*")).bold = True
            else:
                doc.add_paragraph(line)

    add_section("一、通知（正式版）", data.get("notice_formal", ""))
    add_section("二、通知（群发版）", data.get("notice_group", ""))
    add_section("三、通知（家长版）", data.get("notice_parent", ""))
    add_section("四、班会方案", data.get("meeting_plan", ""))
    add_section("五、签到表", data.get("signin_sheet", ""))
    add_section("六、班会纪要模板", data.get("minutes_template", ""))

    disclaimer = doc.add_paragraph()
    run = disclaimer.add_run("⚠️ 本文件由 AI 生成，仅为草稿，涉及学生/思政/危机材料必须人工复核。")
    run.italic = True
    run.font.size = 110000

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
