"""表格台账服务：Excel 读取、批量生成评语、回写导出、信息提取。"""

import io
import json
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook
from dashscope import Generation

from app.core.config import settings
from app.models.doc_task import DocTask
from app.models.operation_log import OperationLog
from app.models.template import Template


BATCH_REVIEW_SYSTEM = """你是一位高校辅导员。请根据学生信息生成个性化的综合评语。
评语要求：
- 150-200字
- 肯定学生的优点和进步
- 温和指出改进方向
- 语气关心、不评判
- 每位学生独立生成，不重复"""

BATCH_REVIEW_USER = """请为以下学生生成综合评语：

学生姓名：{name}
班级：{class_name}
学业表现：{performance}
活动参与：{activities}
特点备注：{notes}

请生成一段个性化评语："""

EXTRACT_SYSTEM = """你是一位细心的学工信息录入助手。请从给出的自由文本中提取学生信息，
填入结构化台账。未提及的字段标记为"【缺失】"，不要编造。"""

EXTRACT_USER = """请从以下文本中提取学生信息，填入台账字段：

## 台账字段
{fields}

## 原始文本
{text}

请返回 JSON 格式，每个字段一个键值对。未提及的字段值为"【缺失】"。只返回 JSON，不返回其他内容。"""


def read_excel(file_content: bytes) -> list[dict]:
    """读取 Excel 文件，返回每行数据的 dict 列表。"""
    wb = load_workbook(io.BytesIO(file_content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel 文件为空")

    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data_rows = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        record = {}
        for i, val in enumerate(row):
            if i < len(headers):
                record[headers[i]] = str(val).strip() if val is not None else ""
        data_rows.append(record)
    return data_rows


REVIEW_ATTEMPTS = 2  # 单条评语最多尝试次数：并发下偶发限流不至于整行作废


def _review_prompt(row: dict, index: int, custom_user_prompt: str | None) -> str:
    """单个学生的评语提示词：有模板用模板，否则用默认模板。"""
    if custom_user_prompt:
        return custom_user_prompt.format(**{k: row.get(k, f"【{k}】") for k in row.keys()})
    return BATCH_REVIEW_USER.format(
        name=row.get("姓名", row.get("学生姓名", f"学生{index+1}")),
        class_name=row.get("班级", row.get("所在班级", "")),
        performance=row.get("学业表现", row.get("成绩", "")),
        activities=row.get("活动参与", row.get("社团活动", "")),
        notes=row.get("备注", row.get("特点", "")),
    )


def _generate_review(row: dict, index: int, custom_user_prompt: str | None) -> str:
    """生成单个学生的评语（纯计算，不碰数据库，可安全并发）。"""
    user_prompt = _review_prompt(row, index, custom_user_prompt)
    last = "生成失败：未知原因"
    for attempt in range(REVIEW_ATTEMPTS):
        try:
            resp = Generation.call(
                model="qwen-plus",
                api_key=settings.dashscope_api_key,
                messages=[
                    {"role": "system", "content": BATCH_REVIEW_SYSTEM},
                    {"role": "user", "content": user_prompt},
                ],
                result_format="message",
            )
            if resp.status_code == 200:
                return resp.output.choices[0].message.content
            last = f"生成失败: {resp.message}"
        except Exception as e:  # noqa: BLE001
            last = f"生成异常: {str(e)[:100]}"
        if attempt + 1 < REVIEW_ATTEMPTS:
            time.sleep(1.5)
    return last


def batch_generate_reviews(
    db: Session,
    user_id: int,
    student_rows: list[dict],
    template_id: int | None = None,
) -> list[dict]:
    """批量生成学生评语。每行一次模型调用，并发执行，结果按原行顺序返回。"""
    if not settings.dashscope_api_key:
        raise RuntimeError("API Key 未配置")

    if not student_rows:
        return []

    # 如果有模板，使用模板的 prompt
    custom_user_prompt = None
    if template_id:
        tmpl = db.query(Template).filter(Template.id == template_id).first()
        if tmpl:
            custom_user_prompt = tmpl.user_prompt_template

    workers = max(1, min(settings.ledger_concurrency, len(student_rows)))
    results: list = [None] * len(student_rows)
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futures = {
            ex.submit(_generate_review, row, i, custom_user_prompt): i
            for i, row in enumerate(student_rows)
        }
        done = 0
        for fut in as_completed(futures):
            i = futures[fut]
            results[i] = {**student_rows[i], "AI评语": fut.result()}
            done += 1
            print(f"  [{done}/{len(student_rows)}] 评语生成完成")

    # 操作日志
    log = OperationLog(
        user_id=user_id,
        action="batch_review",
        detail=f"批量生成 {len(student_rows)} 条评语",
    )
    db.add(log)
    db.commit()

    return results


def write_excel(rows: list[dict], original_headers: list[str] | None = None) -> bytes:
    """将结果写入 Excel 并返回二进制内容。"""
    wb = Workbook()
    ws = wb.active

    # 确定列头
    if original_headers:
        headers = original_headers + ["AI评语"]
    else:
        all_keys = set()
        for row in rows:
            all_keys.update(row.keys())
        headers = sorted(all_keys)
        if "AI评语" in headers:
            headers.remove("AI评语")
            headers.append("AI评语")

    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def extract_to_ledger(text: str, fields: list[str]) -> dict:
    """从自由文本中提取信息写入台账。"""
    if not settings.dashscope_api_key:
        raise RuntimeError("API Key 未配置")

    fields_str = "\n".join(f"- {f}" for f in fields)

    resp = Generation.call(
        model="qwen-plus",
        api_key=settings.dashscope_api_key,
        messages=[
            {"role": "system", "content": EXTRACT_SYSTEM},
            {"role": "user", "content": EXTRACT_USER.format(fields=fields_str, text=text)},
        ],
        result_format="message",
    )

    if resp.status_code != 200:
        raise RuntimeError(f"提取失败: {resp.message}")

    raw = resp.output.choices[0].message.content
    # 尝试解析 JSON
    try:
        # 清理可能的 markdown 包裹
        if raw.startswith("```"):
            raw = raw.split("\n", 1)[1]
            if raw.endswith("```"):
                raw = raw[:-3]
        result = json.loads(raw)
    except json.JSONDecodeError:
        # 解析失败，返回原始文本标记
        result = {"原始输出": raw, "_parse_error": True}

    return result