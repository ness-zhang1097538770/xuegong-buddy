"""学生档案 API。"""

import json
import io
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session
from docx import Document
from app.core.deps import get_db, get_current_user
from app.models.user import User
from app.models.student import Student

router = APIRouter(prefix="/students", tags=["学生档案"])


@router.get("")
def list_students(
    grade: str | None = Query(None),
    class_name: str | None = Query(None),
    risk_level: str | None = Query(None),
    tag: str | None = Query(None),
    search: str | None = Query(None),
    page: int = Query(1, ge=1),
    size: int = Query(50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """学生列表，支持筛选和分页。"""
    query = db.query(Student).filter(Student.owner_id == user.id)

    if grade:
        query = query.filter(Student.grade == grade)
    if class_name:
        query = query.filter(Student.class_name == class_name)
    if tag:
        query = query.filter(Student.tags.contains(tag))
    if search:
        q = f"%{search}%"
        query = query.filter(
            (Student.name.contains(q)) | (Student.student_id.contains(q))
        )
    if risk_level:
        if risk_level == "high":
            query = query.filter(Student.risk_score >= 3)
        elif risk_level == "medium":
            query = query.filter(Student.risk_score.between(1, 2))
        elif risk_level == "none":
            query = query.filter(Student.risk_score == 0)

    # 默认按风险分降序
    total = query.count()
    students = (
        query.order_by(Student.risk_score.desc(), Student.gpa.desc())
        .offset((page - 1) * size)
        .limit(size)
        .all()
    )

    return {
        "total": total,
        "page": page,
        "size": size,
        "students": [
            {
                "id": s.id,
                "student_id": s.student_id,
                "name": s.name,
                "class_name": s.class_name,
                "grade": s.grade,
                "gender": s.gender,
                "dorm_building": s.dorm_building,
                "dorm_room": s.dorm_room,
                "gpa": s.gpa or 0.0,
                "attendance": s.attendance or 100.0,
                "tags": json.loads(s.tags) if s.tags else [],
                "risk_score": s.risk_score or 0,
            }
            for s in students
        ],
    }


@router.get("/filters")
def get_filter_options(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """返回筛选选项：年级、班级列表。"""
    grades = (
        db.query(Student.grade)
        .filter(Student.owner_id == user.id, Student.grade != "")
        .distinct()
        .all()
    )
    classes = (
        db.query(Student.class_name)
        .filter(Student.owner_id == user.id, Student.class_name != "")
        .distinct()
        .all()
    )
    return {
        "grades": sorted([g[0] for g in grades if g[0]]),
        "classes": sorted([c[0] for c in classes if c[0]]),
    }


@router.get("/{student_id}")
def get_student_detail(
    student_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """学生详情（抽屉用）。"""
    s = (
        db.query(Student)
        .filter(Student.id == student_id, Student.owner_id == user.id)
        .first()
    )
    if not s:
        raise HTTPException(status_code=404, detail="学生不存在")

    # 叙事性档案：按时间倒序取该生历次谈话记录
    from app.models.approval import TalkRecord

    talk_records = (
        db.query(TalkRecord)
        .filter(
            TalkRecord.handler_id == user.id,
            TalkRecord.student_id_ref == s.student_id,
        )
        .order_by(TalkRecord.created_at.desc())
        .limit(50)
        .all()
    )

    return {
        "id": s.id,
        "student_id": s.student_id,
        "name": s.name,
        "class_name": s.class_name,
        "grade": s.grade,
        "gender": s.gender,
        "phone": s.phone,
        "dorm_building": s.dorm_building,
        "dorm_room": s.dorm_room,
        "gpa": s.gpa or 0.0,
        "attendance": s.attendance or 100.0,
        "tags": json.loads(s.tags) if s.tags else [],
        "risk_score": s.risk_score or 0,
        "status": s.status,
        "talk_records": [
            {
                "id": t.id,
                "method": t.method,
                "topic": t.topic,
                "content": t.content,
                "conclusion": t.conclusion,
                "mood": t.mood,
                "need_follow": bool(t.need_follow),
                "follow_date": t.follow_date,
                "follow_closed": bool(t.follow_closed),
                "created_at": t.created_at.isoformat() if t.created_at else "",
            }
            for t in talk_records
        ],
    }


@router.get("/{student_id}/handover")
def export_handover_package(
    student_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """换届/转专业交接包：学生叙事性档案导出为 Word。"""
    s = (
        db.query(Student)
        .filter(Student.id == student_id, Student.owner_id == user.id)
        .first()
    )
    if not s:
        raise HTTPException(status_code=404, detail="学生不存在")

    from app.models.approval import TalkRecord

    talks = (
        db.query(TalkRecord)
        .filter(
            TalkRecord.handler_id == user.id,
            TalkRecord.student_id_ref == s.student_id,
        )
        .order_by(TalkRecord.created_at.asc())
        .all()
    )

    doc = Document()
    doc.styles["Normal"].font.name = "SimSun"
    title = doc.add_heading(f"学生交接包 · {s.name}", level=0)
    title.alignment = 1

    doc.add_heading("一、基本信息", level=1)
    tags = "、".join(json.loads(s.tags)) if s.tags else "无"
    info = [
        f"学号：{s.student_id}",
        f"班级：{s.class_name}　年级：{s.grade}",
        f"性别：{s.gender or '未填'}　电话：{s.phone or '未填'}",
        f"宿舍：{s.dorm_building or '未填'} {s.dorm_room or ''}",
        f"GPA：{s.gpa or 0.0}　出勤率：{s.attendance or 100.0}%",
        f"标签：{tags}",
        f"风险分：{s.risk_score or 0}",
    ]
    for line in info:
        doc.add_paragraph(line)

    doc.add_heading("二、谈话叙事档案", level=1)
    if not talks:
        doc.add_paragraph("（暂无谈话记录）")
    for t in talks:
        doc.add_heading(
            f"{t.created_at.strftime('%Y-%m-%d')} · {t.method} · {t.topic or '未归类'}",
            level=2,
        )
        doc.add_paragraph(f"谈话内容：{t.content or '（无）'}")
        doc.add_paragraph(f"结论：{t.conclusion or '（无）'}")
        if t.need_follow:
            doc.add_paragraph(
                f"跟进：{'已闭环' if t.follow_closed else f'待跟进（{t.follow_date}）'}"
            )

    doc.add_heading("三、交接提示", level=1)
    doc.add_paragraph("本交接包由系统根据留存记录生成，供换届/转专业/离校交接使用。请新接手辅导员在正式接手前与学生完成一次面谈，核对信息后更新档案。")

    disclaimer = doc.add_paragraph()
    run = disclaimer.add_run("⚠️ 本文件由系统生成，涉及学生敏感信息，请妥善保管，不得外传。")
    run.italic = True
    run.font.size = 110000

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=handover_{s.student_id}.docx"},
    )


@router.post("/seed-demo")
def seed_demo_students(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """生成演示学生数据（幂等：按学号去重）。"""
    demos = [
        {"student_id": "2023010101", "name": "王雨桐", "class_name": "计算机2301班", "grade": "2023级",
         "gender": "女", "phone": "13800000101", "dorm_building": "紫荆1号楼", "dorm_room": "305",
         "gpa": 3.2, "attendance": 96.0, "tags": json.dumps(["心理关注"], ensure_ascii=False), "risk_score": 2},
        {"student_id": "2023010105", "name": "李亮", "class_name": "计算机2301班", "grade": "2023级",
         "gender": "男", "phone": "13800000105", "dorm_building": "紫荆2号楼", "dorm_room": "402",
         "gpa": 1.8, "attendance": 72.0, "tags": json.dumps(["学业预警"], ensure_ascii=False), "risk_score": 3},
        {"student_id": "2023010108", "name": "李浩然", "class_name": "计算机2301班", "grade": "2023级",
         "gender": "男", "phone": "13900000108", "dorm_building": "紫荆2号楼", "dorm_room": "405",
         "gpa": 1.9, "attendance": 68.0, "tags": json.dumps(["学业预警", "3门挂科"], ensure_ascii=False), "risk_score": 3},
        {"student_id": "2023010112", "name": "张一凡", "class_name": "计算机2301班", "grade": "2023级",
         "gender": "男", "phone": "13800000112", "dorm_building": "紫荆2号楼", "dorm_room": "407",
         "gpa": 3.6, "attendance": 98.0, "tags": json.dumps(["班长"], ensure_ascii=False), "risk_score": 0},
        {"student_id": "2023010115", "name": "陈思远", "class_name": "软件2302班", "grade": "2023级",
         "gender": "男", "phone": "13800000115", "dorm_building": "紫荆3号楼", "dorm_room": "508",
         "gpa": 2.6, "attendance": 82.0, "tags": json.dumps(["兼职"], ensure_ascii=False), "risk_score": 1},
        {"student_id": "2023010122", "name": "赵梦琪", "class_name": "软件2302班", "grade": "2023级",
         "gender": "女", "phone": "13800000122", "dorm_building": "紫荆1号楼", "dorm_room": "310",
         "gpa": 3.4, "attendance": 97.0, "tags": json.dumps(["困难认定"], ensure_ascii=False), "risk_score": 0},
        {"student_id": "2023010130", "name": "刘子墨", "class_name": "软件2302班", "grade": "2023级",
         "gender": "男", "phone": "13800000130", "dorm_building": "紫荆3号楼", "dorm_room": "510",
         "gpa": 2.2, "attendance": 88.0, "tags": json.dumps([]), "risk_score": 1},
        {"student_id": "2023010133", "name": "孙悦", "class_name": "计算机2301班", "grade": "2023级",
         "gender": "女", "phone": "13800000133", "dorm_building": "紫荆1号楼", "dorm_room": "312",
         "gpa": 3.8, "attendance": 99.0, "tags": json.dumps(["学习委员"], ensure_ascii=False), "risk_score": 0},
    ]
    created = 0
    for d in demos:
        exists = db.query(Student).filter(
            Student.student_id == d["student_id"], Student.owner_id == user.id
        ).first()
        if exists:
            continue
        db.add(Student(**d, owner_id=user.id))
        created += 1
    db.commit()
    return {"message": f"已生成 {created} 名演示学生（重复已跳过）"}


@router.post("/import")
async def import_students(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Excel 批量导入学生。"""
    import io

    if not file.filename or not (
        file.filename.endswith(".xlsx") or file.filename.endswith(".xls")
    ):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx / .xls 格式")

    try:
        import openpyxl

        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active

        headers = [cell.value for cell in ws[1]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        imported = 0
        skipped = 0
        errors = []

        for i, row in enumerate(rows):
            if not row or not any(row):
                continue
            data = dict(zip(headers, row))

            student_id = str(data.get("学号", data.get("student_id", ""))).strip()
            name = str(data.get("姓名", data.get("name", ""))).strip()

            if not student_id or not name:
                skipped += 1
                errors.append(f"第{i+2}行：学号或姓名为空，已跳过")
                continue

            # 检查重复
            existing = (
                db.query(Student)
                .filter(
                    Student.student_id == student_id,
                    Student.owner_id == user.id,
                )
                .first()
            )
            if existing:
                # 更新
                existing.name = name
                existing.class_name = str(
                    data.get("班级", data.get("class_name", existing.class_name))
                )
                existing.grade = str(data.get("年级", data.get("grade", "")))
                existing.gender = str(data.get("性别", data.get("gender", "")))
                existing.phone = str(data.get("电话", data.get("phone", "")))
                existing.dorm_building = str(
                    data.get("楼栋", data.get("dorm_building", ""))
                )
                existing.dorm_room = str(data.get("房号", data.get("dorm_room", "")))
                try:
                    existing.gpa = float(data.get("GPA", data.get("gpa", 0)) or 0)
                except (ValueError, TypeError):
                    pass
                try:
                    existing.attendance = float(
                        data.get("出勤率", data.get("attendance", 100)) or 100
                    )
                except (ValueError, TypeError):
                    pass
                tags_raw = str(data.get("标签", data.get("tags", "")))
                if tags_raw:
                    existing.tags = json.dumps(
                        [t.strip() for t in tags_raw.split(",") if t.strip()],
                        ensure_ascii=False,
                    )
                imported += 1
            else:
                tags_raw = str(data.get("标签", data.get("tags", "")))
                tags = (
                    json.dumps(
                        [t.strip() for t in tags_raw.split(",") if t.strip()],
                        ensure_ascii=False,
                    )
                    if tags_raw
                    else "[]"
                )

                student = Student(
                    student_id=student_id,
                    name=name,
                    class_name=str(data.get("班级", data.get("class_name", ""))),
                    grade=str(data.get("年级", data.get("grade", ""))),
                    gender=str(data.get("性别", data.get("gender", ""))),
                    phone=str(data.get("电话", data.get("phone", ""))),
                    dorm_building=str(data.get("楼栋", data.get("dorm_building", ""))),
                    dorm_room=str(data.get("房号", data.get("dorm_room", ""))),
                    gpa=float(data.get("GPA", data.get("gpa", 0)) or 0),
                    attendance=float(
                        data.get("出勤率", data.get("attendance", 100)) or 100
                    ),
                    tags=tags,
                    owner_id=user.id,
                )
                db.add(student)
                imported += 1

        db.commit()
        return {
            "imported": imported,
            "skipped": skipped,
            "total_rows": len(rows),
            "errors": errors[:10],
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导入失败: {str(e)}")